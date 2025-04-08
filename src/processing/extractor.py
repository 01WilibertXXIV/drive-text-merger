import io
import logging
import csv
import openpyxl # For .xlsx
import docx # For .docx
import PyPDF2 # For .pdf
from typing import Optional, Union

# --- Helper Functions for Specific File Types ---

def _extract_text_from_docx(content_bytes: bytes, file_url: str) -> Optional[str]:
    """Extracts text content from DOCX bytes."""
    try:
        document = docx.Document(io.BytesIO(content_bytes))
        full_text = [para.text for para in document.paragraphs]
        return '\n'.join(full_text)
    except Exception as e:
        logging.error(f"Error extracting text from DOCX: {e} (URL: {file_url})", exc_info=True)
        # Return None or a specific error message
        return f"[DOCX text extraction failed. View file at {file_url}]" # Consistent with original behavior


def _extract_text_from_pdf(content_bytes: bytes, file_url: str) -> Optional[str]:
    """Extracts text content from PDF bytes."""
    full_text = []
    try:
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(content_bytes))
        # Handle encrypted PDFs if necessary (optional)
        if pdf_reader.is_encrypted:
            try:
                # Attempt to decrypt with an empty password, common for some PDFs
                pdf_reader.decrypt('')
            except Exception as decrypt_error:
                logging.warning(f"Could not decrypt PDF {file_url}: {decrypt_error}")
                # Decide if you want to return an error or try extraction anyway
                # return f"[PDF is encrypted and could not be decrypted. View file at {file_url}]"

        for page_num, page in enumerate(pdf_reader.pages):
            try:
                page_text = page.extract_text()
                if page_text: # Add text only if extraction returned something
                    full_text.append(page_text)
            except Exception as page_error:
                 logging.warning(f"Error extracting text from PDF page {page_num+1} for {file_url}: {page_error}")
                 # Continue to next page

        if not full_text:
             logging.warning(f"No text could be extracted from PDF {file_url}. It might be image-based or have extraction issues.")
             return f"[No text extracted from PDF. It might be image-based. View file at {file_url}]"

        return '\n'.join(full_text)

    except PyPDF2.errors.PdfReadError as pdf_error:
         logging.error(f"Invalid PDF file {file_url}: {pdf_error}", exc_info=True)
         return f"[Invalid or corrupted PDF file. View file at {file_url}]"
    except Exception as e:
        logging.error(f"Error extracting text from PDF {file_url}: {e}", exc_info=True)
        return f"[PDF text extraction failed. View file at {file_url}]"


def _extract_from_xlsx(content_bytes: bytes, file_url: str) -> Optional[str]:
    """Extracts text content from XLSX bytes."""
    full_text = []
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True) # data_only=True to get values, not formulas
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            # Optional: Add sheet name to the text
            # full_text.append(f"--- Sheet: {sheet_name} ---")
            for row in sheet.iter_rows():
                row_text = []
                for cell in row:
                    if cell.value is not None:
                        # Convert cell value to string, handle various types
                        cell_value_str = str(cell.value).strip()
                        if cell_value_str: # Add only non-empty cells
                             row_text.append(cell_value_str)
                if row_text: # Add row only if it contains text
                     full_text.append(" | ".join(row_text)) # Join cells with a separator
        return '\n'.join(full_text)
    except Exception as e:
        logging.error(f"Error extracting text from XLSX: {e} (URL: {file_url})", exc_info=True)
        return f"[XLSX text extraction failed. View file at {file_url}]"


def _extract_from_csv(content_bytes: bytes, file_url: str) -> Optional[str]:
    """Extracts text content from CSV bytes."""
    full_text = []
    # Try decoding with common encodings
    encodings_to_try = ['utf-8', 'cp1252', 'latin-1']
    decoded_text = None
    for enc in encodings_to_try:
        try:
            decoded_text = content_bytes.decode(enc)
            logging.debug(f"Decoded CSV {file_url} using {enc}")
            break # Success
        except UnicodeDecodeError:
            continue # Try next encoding
        except Exception as decode_error: # Catch other potential errors during decode
             logging.error(f"Error decoding CSV {file_url} with {enc}: {decode_error}")
             return f"[CSV decoding failed. View file at {file_url}]"


    if decoded_text is None:
         logging.error(f"Could not decode CSV {file_url} with attempted encodings.")
         return f"[CSV decoding failed (unknown encoding?). View file at {file_url}]"

    try:
        # Use io.StringIO to treat the decoded string as a file
        csv_file = io.StringIO(decoded_text)
        # Sniff the dialect (delimiter, quote char, etc.) for robustness
        try:
             dialect = csv.Sniffer().sniff(csv_file.read(1024*10)) # Read a sample
             csv_file.seek(0) # Rewind after sniffing
             reader = csv.reader(csv_file, dialect)
        except csv.Error:
             logging.warning(f"Could not automatically detect CSV dialect for {file_url}. Falling back to standard comma delimiter.")
             csv_file.seek(0) # Rewind
             reader = csv.reader(csv_file) # Use default comma delimiter

        for row in reader:
             row_text = [cell.strip() for cell in row if cell and cell.strip()]
             if row_text: # Add row only if it contains text
                 full_text.append(" | ".join(row_text)) # Join cells with a separator
        return '\n'.join(full_text)
    except Exception as e:
        logging.error(f"Error extracting text from CSV: {e} (URL: {file_url})", exc_info=True)
        return f"[CSV text extraction failed. View file at {file_url}]"


# --- Main Dispatcher Function ---

def extract_text(
    content_bytes: bytes,
    downloaded_mime_type: str,
    file_name: str, # Useful for CSV/XLSX context, maybe file type hints
    file_url: str   # Useful for error messages
) -> Optional[str]:
    """
    Extracts text from downloaded file content based on its MIME type.

    Args:
        content_bytes: The raw byte content of the downloaded file.
        downloaded_mime_type: The MIME type of the content AS IT WAS DOWNLOADED
                              (e.g., DOCX even if original was Google Doc).
        file_name: The original name of the file.
        file_url: The web view URL of the file.

    Returns:
        The extracted text as a string, or a placeholder/error string if extraction fails.
        Returns None only in case of unexpected internal error before dispatch.
    """
    extracted_text: Optional[str] = None

    logging.debug(f"Attempting text extraction for '{file_name}' (Type: {downloaded_mime_type})")

    if downloaded_mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
        extracted_text = _extract_text_from_docx(content_bytes, file_url)
    elif downloaded_mime_type == 'application/pdf':
        extracted_text = _extract_text_from_pdf(content_bytes, file_url)
    elif downloaded_mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
        extracted_text = _extract_from_xlsx(content_bytes, file_url)
    elif downloaded_mime_type == 'text/csv':
        extracted_text = _extract_from_csv(content_bytes, file_url)
    else:
        # Should not happen if drive_api.download_file_content only returns supported types,
        # but good to have a fallback.
        logging.warning(f"No specific text extractor available for downloaded mimeType: {downloaded_mime_type} for file {file_name}")
        extracted_text = f"[Unsupported format for text extraction: {downloaded_mime_type}]"

    # Post-processing: Check if extraction actually yielded text
    if extracted_text is not None and not extracted_text.strip():
         # Handle cases where extraction ran but produced empty output (e.g., empty doc)
         # You might return an empty string, None, or a specific note
         logging.info(f"Extraction for '{file_name}' resulted in empty text.")
         # return "" # Return empty string if that's desired for empty files
         # Return the placeholder if it contains one
         if file_url in extracted_text: # Check if it was already an error message
              return extracted_text
         else:
              return "[File appeared empty or text could not be extracted]"


    # Return the extracted text (which might be an error message from the helpers)
    return extracted_text